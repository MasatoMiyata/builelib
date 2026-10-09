"""???????Excel?JSON?????????????"""

import copy
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from builelib import commons
from builelib.input.make_inputdata import make_jsondata_from_Ver2_sheet
from builelib.input.parser import parse_input_sheet
from builelib.input.reference_specification import expand_reference_specifications, lighting_rated_power, prepare_input_data


def _workbook(tmp_path, rooms, lighting_rows):
    workbook = Workbook()
    form1 = workbook.active
    form1.title = "F1"
    form1["A1"] = "F1"
    form1["M2"] = "column bound"
    form4 = workbook.create_sheet("F4")
    for row_number, (name, building, room_type, area) in enumerate(rooms, start=11):
        for column, value in enumerate(("1F", name, building, room_type, area), start=1):
            form1.cell(row_number, column, value)
    for row_number, row in enumerate(lighting_rows, start=11):
        for column, value in row.items():
            form4.cell(row_number, column, value)
    # The legacy parser reads fixed columns through Q, including on continuation rows.
    form4["Q1"] = ""
    form4["Q2"] = "column bound"
    path = tmp_path / "lighting.xlsx"
    workbook.save(path)
    return path


def test_reference_lighting_expands_base_and_continuation_rows(tmp_path):
    path = _workbook(
        tmp_path,
        [("R1", "事務所等", "事務室", 100), ("R2", "事務所等", "事務室", 25)],
        [
            {1: "1F", 2: "R1", 5: 0, 11: "基準設定仕様", 12: 999, 13: 3, 14: "有"},
            {11: "基準設定仕様"},
            {1: "1F", 2: "R2", 5: 25, 11: "基準設定仕様_照明_0001", 12: 10, 13: 2},
        ],
    )
    parsed = parse_input_sheet(path)
    assert not [error for error in parsed.errors if "基準設定仕様" in error]
    units = parsed.data["LightingSystems"]["1F_R1"]["lightingUnit"]
    assert list(units) == ["基準設定仕様_照明_0002", "基準設定仕様_照明_0003"]
    for unit in units.values():
        assert unit == {
            "RatedPower": 1630,
            "Number": 1,
            "OccupantSensingCTRL": "無",
            "IlluminanceSensingCTRL": "無",
            "TimeScheduleCTRL": "無",
            "InitialIlluminationCorrectionCTRL": "無",
        }
    assert parsed.data["LightingSystems"]["1F_R2"]["lightingUnit"]["基準設定仕様_照明_0001"]["RatedPower"] == 10


def test_excel_reader_keeps_reference_request_for_calculation_stage(tmp_path):
    path = _workbook(
        tmp_path,
        [("R1", "事務所等", "事務室", 100)],
        [{1: "1F", 2: "R1", 11: "基準設定仕様"}],
    )

    raw, validation = make_jsondata_from_Ver2_sheet(str(path))
    assert not [error for error in validation["error"] if "基準設定仕様" in error]
    assert not [error for error in validation["error"] if "スキーマエラー" in error]
    assert raw["LightingSystems"]["1F_R1"]["lightingUnit"] == {}
    assert raw["ReferenceSpecificationRequests"] == [
        {"equipment": "照明", "room": "1F_R1", "source": "F4 11行目（機器名称）"}
    ]

    complete, result_validation = prepare_input_data(raw, validation, from_excel=True)
    assert not [error for error in result_validation["error"] if "基準設定仕様" in error]
    assert "ReferenceSpecificationRequests" not in complete
    assert complete["LightingSystems"]["1F_R1"]["lightingUnit"]["基準設定仕様_照明_0001"]["RatedPower"] == 1630


def test_excel_runner_writes_expanded_input(tmp_path):
    from builelib import runner

    path = _workbook(
        tmp_path,
        [("R1", "事務所等", "事務室", 100)],
        [{1: "1F", 2: "R1", 11: "基準設定仕様"}],
    )
    runner.calculate(path, exec_calculation=False)

    output = json.loads((tmp_path / "lighting_input.json").read_text(encoding="utf-8"))
    assert "ReferenceSpecificationRequests" not in output
    assert output["LightingSystems"]["1F_R1"]["lightingUnit"]["基準設定仕様_照明_0001"]["RatedPower"] == 1630


def test_excel_runner_reports_reference_error_with_sheet_row(tmp_path):
    from builelib import runner

    path = _workbook(
        tmp_path,
        [("R1", "共同住宅", "住戸", 100)],
        [{1: "1F", 2: "R1", 11: "基準設定仕様"}],
    )
    runner.calculate(path, exec_calculation=False)

    validation = json.loads((tmp_path / "lighting_validation.json").read_text(encoding="utf-8"))
    assert any("F4 11行目" in error and "共同住宅" in error for error in validation["error"])


def test_reference_lighting_rejects_unsupported_and_bad_area(tmp_path):
    path = _workbook(
        tmp_path,
        [("R1", "共同住宅", "住戸", 30), ("R2", "事務所等", "事務室", 0)],
        [
            {1: "1F", 2: "R1", 5: 30, 11: "基準設定仕様"},
            {1: "1F", 2: "R2", 5: 20, 11: "基準設定仕様"},
        ],
    )
    parsed = parse_input_sheet(path)
    assert any("F4 11行目" in error and "共同住宅" in error and "対象外" in error for error in parsed.errors)
    assert any("F4 12行目" in error and "床面積" in error and "正の数値" in error for error in parsed.errors)
    assert "1F_R1" not in parsed.data["LightingSystems"]
    assert "1F_R2" not in parsed.data["LightingSystems"]


def test_reference_lighting_rejects_missing_database_value(monkeypatch):
    from builelib.input import reference_specification as reference

    source = reference._load_database()
    database = json.loads(json.dumps(source))
    database["ventilation_and_lighting"]["rooms"]["事務所等"]["事務室"]["lighting"]["基準設定消費電力 [W/m2]"] = None
    monkeypatch.setattr(reference, "_load_database", lambda: database)
    try:
        lighting_rated_power("事務所等", "事務室", 100)
    except reference.ReferenceSpecificationError as exc:
        assert "適用できません" in str(exc)
    else:
        raise AssertionError("Missing reference power was accepted")


def test_json_requests_use_same_expansion_without_changing_source():
    source = {
        "Rooms": {
            "1F_R1": {"buildingType": "事務所等", "roomType": "事務室", "roomArea": 100},
        },
        "LightingSystems": {
            "1F_R1": {"roomIndex": 2.5, "lightingUnit": {"基準設定仕様_照明_0001": {"RatedPower": 10}}},
        },
        "ReferenceSpecificationRequests": [
            {"equipment": "照明", "room": "1F_R1"},
            {"equipment": "照明", "room": "1F_R1"},
        ],
    }
    original = copy.deepcopy(source)
    expanded, errors = expand_reference_specifications(source)
    assert errors == []
    assert source == original
    assert "ReferenceSpecificationRequests" not in expanded
    units = expanded["LightingSystems"]["1F_R1"]["lightingUnit"]
    assert list(units) == [
        "基準設定仕様_照明_0001",
        "基準設定仕様_照明_0002",
        "基準設定仕様_照明_0003",
    ]
    assert units["基準設定仕様_照明_0002"]["RatedPower"] == 1630
    assert expanded["LightingSystems"]["1F_R1"]["roomWidth"] is None


def test_json_lighting_request_rejects_area_override():
    source = {
        "Rooms": {
            "1F_R1": {"buildingType": "事務所等", "roomType": "事務室", "roomArea": 100},
        },
        "ReferenceSpecificationRequests": [
            {"equipment": "照明", "room": "1F_R1", "roomArea": 20},
        ],
    }
    _, errors = expand_reference_specifications(source)
    assert any("roomArea" in error for error in errors)


def test_json_requests_report_unsupported_building():
    source = {
        "Rooms": {"1F_R1": {"buildingType": "共同住宅", "roomType": "住戸", "roomArea": 30}},
        "ReferenceSpecificationRequests": [{"equipment": "照明", "room": "1F_R1"}],
    }
    expanded, errors = expand_reference_specifications(source)
    assert any("共同住宅" in error and "対象外" in error for error in errors)
    assert "1F_R1" not in expanded["LightingSystems"]


def test_calculate_from_json_expands_before_lighting_calculation(monkeypatch):
    from builelib import runner

    seen = {}

    def fake_lighting(inputdata, **kwargs):
        seen["units"] = inputdata["LightingSystems"]["1F_R1"]["lightingUnit"]
        return {"for_CGS": {}, "E_lighting": 10, "Es_lighting": 20, "BEI_L": 0.5}

    monkeypatch.setattr(runner.bc, "inputdata_validation", lambda data: [])
    monkeypatch.setattr(runner.lighting, "calc_energy", fake_lighting)
    monkeypatch.setattr(runner.other_energy, "calc_energy", lambda *args, **kwargs: {"for_CGS": {}, "E_other": 0})
    source = {
        "Rooms": {"1F_R1": {"buildingType": "事務所等", "roomType": "事務室", "roomArea": 100}},
        "ReferenceSpecificationRequests": [{"equipment": "照明", "room": "1F_R1"}],
    }
    result = runner.calculate_from_json(source)
    assert result["errors"] == []
    assert seen["units"]["基準設定仕様_照明_0001"]["RatedPower"] == 1630
    assert "LightingSystems" not in source


def test_calculate_from_json_stops_on_reference_error(monkeypatch):
    from builelib import runner

    def must_not_calculate(*args, **kwargs):
        raise AssertionError("入力エラー後に設備計算が呼ばれました")

    monkeypatch.setattr(runner.lighting, "calc_energy", must_not_calculate)
    source = {
        "Rooms": {"1F_R1": {"buildingType": "共同住宅", "roomType": "住戸", "roomArea": 30}},
        "ReferenceSpecificationRequests": [{"equipment": "照明", "room": "1F_R1"}],
    }
    result = runner.calculate_from_json(source)
    assert any("共同住宅" in error for error in result["errors"])


def test_json_file_input_uses_common_expansion(tmp_path, monkeypatch):
    from builelib import runner

    monkeypatch.setattr(runner.bc, "inputdata_validation", lambda data: [])
    source = {
        "Rooms": {"1F_R1": {"buildingType": "事務所等", "roomType": "事務室", "roomArea": 100}},
        "ReferenceSpecificationRequests": [{"equipment": "照明", "room": "1F_R1"}],
    }
    path = tmp_path / "input.json"
    path.write_text(json.dumps(source, ensure_ascii=False), encoding="utf-8")
    runner.calculate(path, exec_calculation=False)
    expanded = json.loads((tmp_path / "input_input.json").read_text(encoding="utf-8"))
    validation = json.loads((tmp_path / "input_validation.json").read_text(encoding="utf-8"))
    assert validation["error"] == []
    assert expanded["LightingSystems"]["1F_R1"]["lightingUnit"]["基準設定仕様_照明_0001"]["RatedPower"] == 1630
    assert "ReferenceSpecificationRequests" not in expanded


def test_expanded_real_json_passes_input_schema():
    from builelib import commons

    example = Path(__file__).parent / "whole_building" / "sample01_WEBPRO_inputSheet_for_Ver3.8_input.json"
    source = json.loads(example.read_text(encoding="utf-8"))
    source["ReferenceSpecificationRequests"] = [{"equipment": "照明", "room": "1F_ロビー"}]
    expanded, errors = expand_reference_specifications(source)
    assert errors == []
    assert commons.inputdata_validation(expanded) == []


SAMPLE_JSON = Path(__file__).parent / "whole_building" / "sample01_WEBPRO_inputSheet_for_Ver3.8_input.json"
SAMPLE_XLSX = Path(__file__).parent / "whole_building" / "sample01_WEBPRO_inputSheet_for_Ver3.8.xlsx"


def _sample_with_requests():
    source = json.loads(SAMPLE_JSON.read_text(encoding="utf-8"))
    wall = source["EnvelopeSet"]["1F_ロビー"]["WallList"][0]
    wall["WallSpec"] = "基準設定仕様"
    wall["WindowList"][0]["WindowID"] = "基準設定仕様"
    source["ReferenceSpecificationRequests"] = [
        {"equipment": "外壁", "zone": "1F_ロビー", "wallIndex": 0},
        {"equipment": "窓", "zone": "1F_ロビー", "wallIndex": 0, "windowIndex": 0},
        {"equipment": "換気", "room": "1F_便所1", "unitType": "排気"},
        {"equipment": "給湯", "room": "1F_中央監視室・警備室", "savingSystem": "無"},
        {"equipment": "昇降機", "room": "7F_事務室1"},
    ]
    return source


def test_all_five_non_lighting_equipment_expand_to_valid_calculation_json():
    source = _sample_with_requests()
    expanded, errors = expand_reference_specifications(source)
    assert errors == []
    assert commons.inputdata_validation(expanded) == []
    assert "ReferenceSpecificationRequests" in source  # 元のJSONは変更しない。
    assert "ReferenceSpecificationRequests" not in expanded

    wall = expanded["EnvelopeSet"]["1F_ロビー"]["WallList"][0]
    assert wall["WallSpec"] == "基準設定仕様_外壁_0001"
    assert expanded["WallConfigure"][wall["WallSpec"]]["Uvalue"] > 0
    window = wall["WindowList"][0]
    assert window["WindowID"] == "基準設定仕様_窓_0001"
    assert window["isBlind"] == "無"
    assert expanded["WindowConfigure"][window["WindowID"]]["windowIvalue"] > 0

    fan = expanded["VentilationUnit"]["基準設定仕様_換気_0001"]
    assert fan["FanAirVolume"] == 40.5 * 33.28
    assert fan["MoterRatedPower"] == 0.010125 * 33.28
    assert expanded["VentilationRoom"]["1F_便所1"]["VentilationUnitRef"]["基準設定仕様_換気_0001"]["UnitType"] == "排気"

    hotwater = expanded["HotwaterSupplySystems"]["基準設定仕様_給湯_0001"]
    assert hotwater["HeatSourceUnit"][0]["RatedCapacity"] == 10
    assert hotwater["HeatSourceUnit"][0]["RatedFuelConsumption"] == 12.5
    assert hotwater["PipeSize"] == 40
    assert expanded["HotwaterRoom"]["1F_中央監視室・警備室"]["HotwaterSystem"][-1]["SystemName"] == "基準設定仕様_給湯_0001"

    elevator = expanded["Elevators"]["7F_事務室1"]["Elevator"][-1]
    assert elevator["ElevatorName"] == "基準設定仕様_昇降機_0001"
    assert (elevator["Number"], elevator["LoadLimit"], elevator["Velocity"]) == (2, 1150, 120)


def test_excel_markers_use_common_expansion(tmp_path):
    workbook = load_workbook(SAMPLE_XLSX)
    workbook["2-4) 外皮 "]["F11"] = "基準設定仕様"
    workbook["2-4) 外皮 "]["H11"] = "基準設定仕様"
    workbook["3-1) 換気室"]["G11"] = "基準設定仕様"
    workbook["3-1) 換気室"]["E11"] = 0  # 様式3-1の床面積は使用しない。
    workbook["5-1) 給湯室"]["H11"] = "基準設定仕様"
    elevator_sheet = workbook["6) 昇降機"]
    elevator_sheet["E11"] = "基準設定仕様"
    for column in "FGHIJ":
        elevator_sheet[f"{column}11"] = None
    path = tmp_path / "reference.xlsx"
    workbook.save(path)
    from builelib.input.make_inputdata import make_jsondata_from_Ver2_sheet

    raw, excel_validation = make_jsondata_from_Ver2_sheet(str(path))
    assert not [error for error in excel_validation["error"] if "基準設定仕様" in error]
    assert raw["EnvelopeSet"]["1F_ロビー"]["WallList"][0]["WallSpec"] == "基準設定仕様"
    raw_reference_elevators = [
        elevator
        for room in raw["Elevators"].values()
        for elevator in room["Elevator"]
        if elevator["ElevatorName"] == "基準設定仕様"
    ]
    assert len(raw_reference_elevators) == 1
    assert raw_reference_elevators[0]["Number"] is None
    assert {request["equipment"] for request in raw["ReferenceSpecificationRequests"]} >= {
        "外壁", "窓", "換気", "給湯", "昇降機"
    }
    for request in raw["ReferenceSpecificationRequests"]:
        assert "11行目" in request["source"]

    parsed = parse_input_sheet(path)
    assert parsed.errors == []
    wall = parsed.data["EnvelopeSet"]["1F_ロビー"]["WallList"][0]
    assert wall["WallSpec"].startswith("基準設定仕様_外壁_")
    assert wall["WindowList"][0]["WindowID"].startswith("基準設定仕様_窓_")
    assert "基準設定仕様_換気_0001" in parsed.data["VentilationUnit"]
    fan = parsed.data["VentilationUnit"]["基準設定仕様_換気_0001"]
    room_area = parsed.data["Rooms"]["1F_便所1"]["roomArea"]
    assert fan["FanAirVolume"] == 40.5 * room_area
    assert fan["MoterRatedPower"] == 0.010125 * room_area
    assert "基準設定仕様_給湯_0001" in parsed.data["HotwaterSupplySystems"]
    assert any(
        elevator["ElevatorName"].startswith("基準設定仕様_昇降機_")
        for room in parsed.data["Elevators"].values()
        for elevator in room["Elevator"]
    )


@pytest.mark.parametrize("window_name", ["G1", "基準設定仕様"])
def test_excel_window_area_checked_during_sheet_reading(tmp_path, window_name):
    from builelib.input.make_inputdata import make_jsondata_from_Ver2_sheet

    workbook = load_workbook(SAMPLE_XLSX)
    sheet = workbook["2-4) 外皮 "]
    sheet["G11"] = 10
    sheet["H11"] = window_name
    path = tmp_path / "window_area.xlsx"
    workbook.save(path)

    _, validation = make_jsondata_from_Ver2_sheet(str(path))
    area_errors = [error for error in validation["error"] if "窓面積が外皮面積よりも大き" in error]
    assert len(area_errors) == 1

    if window_name == "基準設定仕様":
        parsed = parse_input_sheet(path)
        assert len([error for error in parsed.errors if "窓面積が外皮面積よりも大き" in error]) == 1


def test_horizontal_wall_uses_roof_value_and_skips_existing_name():
    from builelib.input.reference_specification import _load_database

    source = json.loads(SAMPLE_JSON.read_text(encoding="utf-8"))
    wall = source["EnvelopeSet"]["1F_ロビー"]["WallList"][0]
    wall["Direction"] = "水平（下）"
    wall["WallSpec"] = "基準設定仕様"
    source["WallConfigure"]["基準設定仕様_外壁_0001"] = copy.deepcopy(source["WallConfigure"]["W1"])
    source["ReferenceSpecificationRequests"] = [{"equipment": "外壁", "zone": "1F_ロビー", "wallIndex": 0}]
    expanded, errors = expand_reference_specifications(source)
    assert errors == []
    assert expanded["EnvelopeSet"]["1F_ロビー"]["WallList"][0]["WallSpec"] == "基準設定仕様_外壁_0002"
    db = _load_database()
    region = int(source["Building"]["Region"])
    sheet = next(sheet for sheet in db["air_conditioning_by_region"].values() if region in sheet["regions"])
    expected = sheet["rooms"]["事務所等"]["ロビー"]["熱貫流率(屋根)"]
    assert expanded["WallConfigure"]["基準設定仕様_外壁_0002"]["Uvalue"] == expected


def test_non_applicable_ventilation_is_an_input_error():
    source = json.loads(SAMPLE_JSON.read_text(encoding="utf-8"))
    source["ReferenceSpecificationRequests"] = [{"equipment": "換気", "room": "1F_ロビー", "unitType": "排気"}]
    expanded, errors = expand_reference_specifications(source)
    assert any("換気仕様を適用できません" in error for error in errors)
    assert "基準設定仕様_換気_0001" not in expanded["VentilationUnit"]


def test_ventilation_request_uses_only_rooms_area():
    source = json.loads(SAMPLE_JSON.read_text(encoding="utf-8"))
    request = {"equipment": "換気", "room": "1F_便所1", "unitType": "排気"}
    source["ReferenceSpecificationRequests"] = [{**request, "roomArea": 20}]
    _, errors = expand_reference_specifications(source)
    assert any("roomArea" in error for error in errors)

    source["ReferenceSpecificationRequests"] = [request]
    source["Rooms"]["1F_便所1"]["roomArea"] = 0
    expanded, errors = expand_reference_specifications(source)
    assert any("床面積" in error and "正の数値" in error for error in errors)
    assert "基準設定仕様_換気_0001" not in expanded["VentilationUnit"]


def test_region_eight_hotwater_efficiency_and_joint_housing_exclusion():
    source = json.loads(SAMPLE_JSON.read_text(encoding="utf-8"))
    source["Building"]["Region"] = "8"
    source["ReferenceSpecificationRequests"] = [
        {"equipment": "給湯", "room": "1F_中央監視室・警備室", "savingSystem": "無"}
    ]
    expanded, errors = expand_reference_specifications(source)
    assert errors == []
    fuel = expanded["HotwaterSupplySystems"]["基準設定仕様_給湯_0001"]["HeatSourceUnit"][0]["RatedFuelConsumption"]
    assert fuel == 10 / 0.82

    source["Rooms"]["1F_中央監視室・警備室"]["buildingType"] = "共同住宅"
    _, errors = expand_reference_specifications(source)
    assert any("共同住宅" in error and "対象外" in error for error in errors)


def test_missing_window_database_value_reports_source_row(monkeypatch):
    from builelib.input import reference_specification as reference

    database = copy.deepcopy(reference._load_database())
    source = json.loads(SAMPLE_JSON.read_text(encoding="utf-8"))
    region = int(source["Building"]["Region"])
    sheet = next(sheet for sheet in database["air_conditioning_by_region"].values() if region in sheet["regions"])
    record = sheet["rooms"]["事務所等"]["ロビー"]
    record["熱貫流率(窓)"] = None
    monkeypatch.setattr(reference, "_load_database", lambda: database)
    window = source["EnvelopeSet"]["1F_ロビー"]["WallList"][0]["WindowList"][0]
    window["WindowID"] = "基準設定仕様"
    source["ReferenceSpecificationRequests"] = [
        {"equipment": "窓", "zone": "1F_ロビー", "wallIndex": 0, "windowIndex": 0}
    ]
    expanded, errors = expand_reference_specifications(source)
    assert any("熱貫流率(窓)" in error and f"{record['_source_row']}行目" in error for error in errors)
    assert "基準設定仕様_窓_0001" not in expanded["WindowConfigure"]


def test_malformed_json_request_returns_validation_error():
    source = _sample_with_requests()
    source["Rooms"]["1F_ロビー"]["buildingType"] = []
    _, errors = expand_reference_specifications(source)
    assert any("建物用途・室用途が不正" in error for error in errors)

# ---------------------------------------------------------------------------
# 空調の基準設定仕様（仕様書第2章）
# ---------------------------------------------------------------------------
AC_SAMPLE_XLSX = Path(__file__).resolve().parents[1] / "examples" / "WEBPRO" / "sample01_WEBPRO_inputSheet_for_Ver3.8.xlsx"
AC_MARKER = "基準設定仕様"


def _ac_zone_input(building="事務所等", room_type="事務室", area=100, region="6"):
    return {
        "Building": {"Region": region},
        "Rooms": {"Z": {"buildingType": building, "roomType": room_type, "roomArea": area}},
        "AirConditioningZone": {"Z": {
            "AHU_cooling_insideLoad": AC_MARKER, "AHU_cooling_outdoorLoad": AC_MARKER,
            "AHU_heating_insideLoad": AC_MARKER, "AHU_heating_outdoorLoad": AC_MARKER,
        }},
        "ReferenceSpecificationRequests": [{"equipment": "空調機群", "zone": "Z"}],
    }


def test_one_ahu_generates_heat_sources_and_pumps():
    source = _ac_zone_input()
    original = copy.deepcopy(source)
    data, errors = expand_reference_specifications(source)
    assert errors == []
    assert source == original
    assert "ReferenceSpecificationRequests" not in data

    zone = data["AirConditioningZone"]["Z"]
    assert zone["AHU_cooling_insideLoad"] == zone["AHU_cooling_outdoorLoad"]
    ahu = data["AirHandlingSystem"][zone["AHU_cooling_insideLoad"]]
    assert ahu["AirHandlingUnit"][0]["RatedCapacityCooling"] == pytest.approx(12)
    assert ahu["AirHandlingUnit"][0]["FanAirVolume"] == pytest.approx(0.3 * 1000 * 100 * 0.0216)
    assert ahu["AirHandlingUnit"][0]["isAirHeatExchanger"] == "全熱交換器あり・様式2-9記載無し"
    cooling = data["HeatsourceSystem"][ahu["HeatSource_cooling"]]["冷房"]["Heatsource"]
    assert len(cooling) == 2
    assert cooling[0]["HeatsourceRatedCapacity"] == pytest.approx(7.3)
    pumps = data["SecondaryPumpSystem"][ahu["Pump_cooling"]]["冷房"]["SecondaryPump"]
    assert len(pumps) == 2
    assert sum(p["RatedPowerConsumption"] for p in pumps) == pytest.approx(14.6 / 22)


def test_two_ahu_groups_use_room_and_outdoor_area_factors():
    data, errors = expand_reference_specifications(_ac_zone_input("ホテル等", "客室"))
    assert errors == []
    zone = data["AirConditioningZone"]["Z"]
    inside = data["AirHandlingSystem"][zone["AHU_cooling_insideLoad"]]
    outside = data["AirHandlingSystem"][zone["AHU_cooling_outdoorLoad"]]
    assert inside is not outside
    assert outside["AirHandlingUnit"][0]["Type"] == "空調機"
    assert inside["AirHandlingUnit"][0]["FanAirVolume"] is None
    assert outside["AirHandlingUnit"][0]["FanAirVolume"] is None
    assert inside["AirHandlingUnit"][0]["RatedCapacityCooling"] == pytest.approx(6)
    assert outside["AirHandlingUnit"][0]["RatedCapacityCooling"] == pytest.approx(3)
    inside_source = data["HeatsourceSystem"][inside["HeatSource_cooling"]]["冷房"]["Heatsource"][0]
    outside_source = data["HeatsourceSystem"][outside["HeatSource_cooling"]]["冷房"]["Heatsource"][0]
    assert inside_source["HeatsourceRatedCapacity"] == pytest.approx(0.054 * 100 * 2 / 3)
    assert outside_source["HeatsourceRatedCapacity"] == pytest.approx(0.054 * 100 / 3)


def test_shared_ahu_uses_representative_usage_and_four_pumps():
    source = _ac_zone_input("物販店舗等", "大型店の売場", area=100)
    data, errors = expand_reference_specifications(source)
    assert errors == []
    ahu = next(iter(data["AirHandlingSystem"].values()))
    pumps = data["SecondaryPumpSystem"][ahu["Pump_cooling"]]["冷房"]["SecondaryPump"]
    assert len(pumps) == 4

    # 室負荷20m²と外気負荷20m²で同点なら、H28表の行が早い事務室を代表にする。
    shared = {
        "Building": {"Region": "6"},
        "Rooms": {
            "office": {"buildingType": "事務所等", "roomType": "事務室", "roomArea": 30},
            "hotel": {"buildingType": "ホテル等", "roomType": "客室", "roomArea": 60},
        },
        "AirConditioningZone": {
            "office": {"AHU_cooling_insideLoad": "shared", "AHU_cooling_outdoorLoad": "other"},
            "hotel": {"AHU_cooling_insideLoad": "other", "AHU_cooling_outdoorLoad": "shared"},
        },
        "AirHandlingSystem": {"shared": {"HeatSource_cooling": AC_MARKER}},
        "ReferenceSpecificationRequests": [{"equipment": "熱源群", "ahu": "shared", "mode": "冷房"}],
    }
    expanded, errors = expand_reference_specifications(shared)
    assert errors == []
    name = expanded["AirHandlingSystem"]["shared"]["HeatSource_cooling"]
    assert expanded["HeatsourceSystem"][name]["冷房"]["Heatsource"][0]["HeatsourceRatedCapacity"] == pytest.approx(0.073 * 40)


def test_no_pump_and_partial_zone_marker():
    no_pump, errors = expand_reference_specifications(
        _ac_zone_input("学校等", "小中学校の教室", region="3")
    )
    assert errors == []
    ahu = next(iter(no_pump["AirHandlingSystem"].values()))
    assert ahu["Pump_cooling"] is None
    assert ahu["Pump_heating"] is None
    assert no_pump.get("SecondaryPumpSystem", {}) == {}

    partial = _ac_zone_input()
    partial["AirConditioningZone"]["Z"]["AHU_cooling_outdoorLoad"] = "existing"
    _, errors = expand_reference_specifications(partial)
    assert any("両方" in error for error in errors)


@pytest.mark.parametrize(
    ("building", "room_type"),
    [
        ("ホテル等", "客室"),
        ("ホテル等", "客室内の浴室等"),
        ("学校等", "宿直室"),
    ],
)
def test_region8_second_ahu_uses_corrected_fan_control(building, room_type):
    # 更新されたH28表のBM13・BM14・BM77にある2台目の制御方式を検証する。
    data, errors = expand_reference_specifications(_ac_zone_input(building, room_type, region="8"))
    assert errors == []
    zone = data["AirConditioningZone"]["Z"]
    assert zone["AHU_cooling_insideLoad"] != zone["AHU_cooling_outdoorLoad"]
    outside = data["AirHandlingSystem"][zone["AHU_cooling_outdoorLoad"]]
    assert outside["AirHandlingUnit"][0]["FanControlType"] == "定風量制御"


def test_excel_collects_requests_and_prepares_all_ac_equipment(tmp_path):
    workbook = load_workbook(AC_SAMPLE_XLSX)
    workbook["2-1) 空調ゾーン"]["J11"] = AC_MARKER
    workbook["2-1) 空調ゾーン"]["K11"] = AC_MARKER
    for column in "VWXY":
        workbook["2-7) 空調機"][f"{column}27"] = AC_MARKER
    path = tmp_path / "ac_reference.xlsx"
    workbook.save(path)

    raw, excel_validation = make_jsondata_from_Ver2_sheet(str(path))
    assert excel_validation["error"] == []
    requests = raw["ReferenceSpecificationRequests"]
    assert {request["equipment"] for request in requests} == {"空調機群", "熱源群", "二次ポンプ群"}
    assert any("2-1) 空調ゾーン 11行目" in request["source"] for request in requests)
    assert any("2-7) 空調機 27行目" in request["source"] for request in requests)

    complete, validation = prepare_input_data(raw, excel_validation, from_excel=True)
    assert validation["error"] == []
    assert commons.inputdata_validation(complete) == []
    assert "ReferenceSpecificationRequests" not in complete
    assert complete["AirConditioningZone"]["1F_ロビー"]["AHU_cooling_insideLoad"].startswith("基準設定仕様_空調機群_")
    assert complete["AirHandlingSystem"]["HU-11"]["HeatSource_cooling"].startswith("基準設定仕様_熱源群_")

    # 片方だけを指定した場合、展開段階のエラーにもExcelのシート・行が残る。
    partial = copy.deepcopy(raw)
    partial["AirConditioningZone"]["1F_ロビー"]["AHU_cooling_outdoorLoad"] = "HU-11"
    _, partial_validation = prepare_input_data(partial, excel_validation, from_excel=True)
    assert any("2-1) 空調ゾーン 11行目" in error and "両方" in error for error in partial_validation["error"])


@pytest.mark.parametrize("region", ("6", "8"))
def test_two_ahu_json_runs_calculation(region):
    from builelib.runner import calculate_from_json

    source = json.loads(SAMPLE_JSON.read_text(encoding="utf-8"))
    source["Building"]["Region"] = region
    zone_name = "1F_ロビー"
    source["Rooms"][zone_name]["buildingType"] = "ホテル等"
    source["Rooms"][zone_name]["roomType"] = "客室"
    for key in (
        "AHU_cooling_insideLoad", "AHU_cooling_outdoorLoad",
        "AHU_heating_insideLoad", "AHU_heating_outdoorLoad",
    ):
        source["AirConditioningZone"][zone_name][key] = AC_MARKER
    source["ReferenceSpecificationRequests"] = [{"equipment": "空調機群", "zone": zone_name}]
    result = calculate_from_json(source)
    assert result["errors"] == []
    assert isinstance(result["result"]["BEI_AC"], float)
